"""Smoke tests for utils.py — pure helpers and registry integrity, no live Azure."""

import enum
import json
import re
import shutil
from pathlib import Path

import pytest

import utils


def test_get_version_returns_nonempty_string():
    version = utils.get_version()
    assert isinstance(version, str)
    assert version


def test_get_current_timestamp_is_mm_dd_yyyy():
    assert re.fullmatch(r"\d{2}\.\d{2}\.\d{4}", utils.get_current_timestamp())


def test_is_auto_run_reads_env(monkeypatch):
    monkeypatch.delenv("STRATUSSCAN_AUTO_RUN", raising=False)
    assert utils.is_auto_run() is False
    monkeypatch.setenv("STRATUSSCAN_AUTO_RUN", "1")
    assert utils.is_auto_run() is True
    monkeypatch.setenv("STRATUSSCAN_AUTO_RUN", "0")
    assert utils.is_auto_run() is False


def test_get_auto_subscriptions_parses_csv(monkeypatch):
    monkeypatch.delenv("STRATUSSCAN_SUBSCRIPTIONS", raising=False)
    assert utils.get_auto_subscriptions() == []
    monkeypatch.setenv("STRATUSSCAN_SUBSCRIPTIONS", " sub-a , sub-b ,, sub-c ")
    assert utils.get_auto_subscriptions() == ["sub-a", "sub-b", "sub-c"]


def test_create_export_filename_shape():
    path = utils.create_export_filename("My Sub Prod", "virtual-machines", "all")
    name = path.replace("\\", "/").split("/")[-1]
    assert name.endswith(".xlsx")
    assert name.startswith("MY-SUB-PROD-")
    assert "-virtual-machines-all-export-" in name
    assert re.search(r"-export-\d{2}\.\d{2}\.\d{4}\.xlsx$", name)


def test_create_export_filename_uppercases_and_sanitizes():
    name = utils.create_export_filename("a/b:c*d", "rg", "all").replace("\\", "/").split("/")[-1]
    # path-unsafe characters collapse to hyphens, result is upper-cased
    assert name.startswith("A-B-C-D-")


def test_archive_outputs_includes_optional_label(monkeypatch, tmp_path):
    output_dir = tmp_path / "output"
    output_dir.mkdir()
    (output_dir / "sample.xlsx").write_text("xlsx")
    monkeypatch.setattr(utils, "__file__", str(tmp_path / "utils.py"))
    monkeypatch.setattr(utils, "get_current_timestamp", lambda: "06.15.2026")

    zip_path = utils.archive_outputs("tier1")

    assert zip_path is not None
    assert zip_path.replace("\\", "/").endswith("/output/exports-tier1-06.15.2026.zip")


def test_extract_resource_group_handles_casing_and_missing():
    assert (
        utils.extract_resource_group(
            "/subscriptions/s/resourceGroups/Prod-RG/providers/Microsoft.Compute/x/y"
        )
        == "Prod-RG"
    )
    assert (
        utils.extract_resource_group(
            "/subscriptions/s/RESOURCEGROUPS/lower/providers/p/x/y"
        )
        == "lower"
    )
    assert utils.extract_resource_group("") == ""
    assert utils.extract_resource_group(None) == ""


def test_detect_environment_defaults_to_public(monkeypatch, tmp_path):
    monkeypatch.delenv("AZURE_ENVIRONMENT", raising=False)
    monkeypatch.setenv("AZURE_CONFIG_DIR", str(tmp_path))
    monkeypatch.setattr(utils, "get_config", lambda: {})
    assert utils.detect_environment() == "public"


def test_detect_environment_government_from_env(monkeypatch):
    monkeypatch.setenv("AZURE_ENVIRONMENT", "AzureUSGovernment")
    monkeypatch.setattr(utils, "get_config", lambda: {})
    assert utils.detect_environment() == "government"


def test_detect_environment_from_config_when_env_absent(monkeypatch):
    monkeypatch.delenv("AZURE_ENVIRONMENT", raising=False)
    monkeypatch.setattr(utils, "get_config", lambda: {"environment": "government"})
    assert utils.detect_environment() == "government"


def test_detect_azure_cloud_reads_az_config(monkeypatch, tmp_path):
    (tmp_path / "config").write_text("[cloud]\nname = AzureUSGovernment\n")
    monkeypatch.setenv("AZURE_CONFIG_DIR", str(tmp_path))
    assert utils.detect_azure_cloud() == "government"

    (tmp_path / "config").write_text("[cloud]\nname = AzureCloud\n")
    assert utils.detect_azure_cloud() == "public"


def test_detect_azure_cloud_returns_none_when_absent(monkeypatch, tmp_path):
    monkeypatch.setenv("AZURE_CONFIG_DIR", str(tmp_path))  # empty dir, no config file
    assert utils.detect_azure_cloud() is None


@pytest.fixture
def isolated_config(monkeypatch, tmp_path):
    """Real config.json and az CLI config files under tmp_path; nothing read from the repo or $HOME."""
    config_path = tmp_path / "config.json"
    az_dir = tmp_path / "azure"
    az_dir.mkdir()
    monkeypatch.setattr(utils, "_CONFIG_PATH", config_path)
    monkeypatch.setattr(utils, "_config_cache", None)
    monkeypatch.setenv("AZURE_CONFIG_DIR", str(az_dir))
    # setenv first so teardown restores the variable even if the code under test writes it
    monkeypatch.setenv("AZURE_ENVIRONMENT", "")
    monkeypatch.delenv("AZURE_ENVIRONMENT")

    def write(config=None, az_cloud=None):
        if config is not None:
            config_path.write_text(json.dumps(config), encoding="utf-8")
        if az_cloud is not None:
            (az_dir / "config").write_text(f"[cloud]\nname = {az_cloud}\n", encoding="utf-8")
        monkeypatch.setattr(utils, "_config_cache", None)

    write.config_path = config_path
    return write


_TEMPLATE = Path(utils.__file__).parent / "config-template.json"


def test_shipped_template_is_auto_and_carries_no_subscriptions():
    template = json.loads(_TEMPLATE.read_text(encoding="utf-8"))
    assert template["environment"] == "auto"
    assert template["subscriptions"] == []
    assert template["default_subscription_id"] == ""
    assert set(template) == set(utils._DEFAULT_CONFIG)


@pytest.mark.parametrize(
    "az_cloud, expected",
    [("AzureUSGovernment", "government"), ("AzureCloud", "public"), (None, "public")],
)
def test_shipped_template_config_reaches_autodetect(isolated_config, az_cloud, expected):
    shutil.copy(_TEMPLATE, isolated_config.config_path)
    isolated_config(az_cloud=az_cloud)
    assert utils.detect_environment() == expected


def test_absent_config_json_reaches_autodetect(isolated_config):
    isolated_config(az_cloud="AzureUSGovernment")
    assert not isolated_config.config_path.exists()
    assert utils.detect_environment() == "government"


@pytest.mark.parametrize("value", [None, "", "  ", "auto", "AUTO", 7, ["government"], "not-a-cloud"])
def test_non_explicit_config_environment_falls_through_to_autodetect(isolated_config, value):
    isolated_config(config={"environment": value}, az_cloud="AzureUSGovernment")
    assert utils.detect_environment() == "government"


def test_config_without_environment_key_falls_through_to_autodetect(isolated_config):
    isolated_config(config={"subscriptions": []}, az_cloud="AzureUSGovernment")
    assert utils.detect_environment() == "government"


def test_corrupt_config_json_falls_through_to_autodetect(isolated_config):
    isolated_config.config_path.write_text("{not json", encoding="utf-8")
    isolated_config(az_cloud="AzureUSGovernment")
    assert utils.detect_environment() == "government"


@pytest.mark.parametrize(
    "configured, az_cloud, expected",
    [
        ("public", "AzureUSGovernment", "public"),
        ("AzurePublicCloud", "AzureUSGovernment", "public"),
        ("government", "AzureCloud", "government"),
        ("AzureUSGovernment", "AzureCloud", "government"),
    ],
)
def test_explicit_config_environment_wins_over_autodetect(isolated_config, configured, az_cloud, expected):
    isolated_config(config={"environment": configured}, az_cloud=az_cloud)
    assert utils.detect_environment() == expected


def test_environment_variable_wins_over_config_and_autodetect(isolated_config, monkeypatch):
    isolated_config(config={"environment": "public"}, az_cloud="AzureCloud")
    monkeypatch.setenv("AZURE_ENVIRONMENT", "AzureUSGovernment")
    assert utils.detect_environment() == "government"

    isolated_config(config={"environment": "government"}, az_cloud="AzureUSGovernment")
    monkeypatch.setenv("AZURE_ENVIRONMENT", "public")
    assert utils.detect_environment() == "public"


@pytest.mark.parametrize("value", ["AzureChinaCloud", "gov", "auto", "AzureUSGovernmentCloud"])
def test_unrecognized_environment_variable_fails_closed(isolated_config, monkeypatch, value):
    isolated_config(config={"environment": "government"}, az_cloud="AzureUSGovernment")
    monkeypatch.setenv("AZURE_ENVIRONMENT", value)
    with pytest.raises(ValueError) as excinfo:
        utils.detect_environment()
    assert value in str(excinfo.value)
    assert "azureusgovernment" in str(excinfo.value)
    assert "azurepubliccloud" in str(excinfo.value)


def test_blank_environment_variable_is_treated_as_unset(isolated_config, monkeypatch):
    isolated_config(az_cloud="AzureUSGovernment")
    monkeypatch.setenv("AZURE_ENVIRONMENT", "   ")
    assert utils.detect_environment() == "government"


def test_save_config_creates_config_json_when_absent(isolated_config):
    assert not isolated_config.config_path.exists()
    utils.save_config({**utils._DEFAULT_CONFIG, "environment": "government"})
    assert json.loads(isolated_config.config_path.read_text(encoding="utf-8"))["environment"] == "government"
    assert utils.detect_environment() == "government"


def test_get_config_defaults_to_auto_when_config_json_absent(isolated_config):
    isolated_config()
    assert utils.get_config()["environment"] == "auto"


_CHAINED_AUTH_MESSAGE = (
    "DefaultAzureCredential failed to retrieve a token from the included credentials.\n"
    "Attempted credentials:\n"
    "\tEnvironmentCredential: EnvironmentCredential authentication unavailable.\n"
    "\tManagedIdentityCredential: (AudienceNotSupported) Audience https://management.azure.com "
    "is not a supported MSI token audience.\n"
    "To mitigate this issue, please refer to the troubleshooting guidelines."
)


def _subscription_client_raising(error):
    class _Subscriptions:
        def list(self):
            raise error
            yield  # pragma: no cover — makes this a lazy pager like the SDK's ItemPaged

    class _Client:
        subscriptions = _Subscriptions()

    return _Client()


def test_list_subscriptions_raises_on_auth_failure_instead_of_returning_empty(monkeypatch):
    exceptions = pytest.importorskip("azure.core.exceptions")
    error = exceptions.ClientAuthenticationError(message=_CHAINED_AUTH_MESSAGE)
    monkeypatch.setattr(utils, "get_azure_client", lambda service: _subscription_client_raising(error))
    monkeypatch.setattr(utils, "detect_environment", lambda: "public")

    with pytest.raises(utils.AzureAccessError) as excinfo:
        utils.list_subscriptions()

    assert excinfo.value.__cause__ is error
    assert excinfo.value.environment == "public"
    assert "ClientAuthenticationError" in str(excinfo.value)
    assert "AudienceNotSupported" in str(excinfo.value)
    assert "\n" not in str(excinfo.value)
    assert "AzurePublicCloud" in excinfo.value.hint
    assert "AZURE_ENVIRONMENT" in excinfo.value.hint


@pytest.mark.parametrize("error_name", ["HttpResponseError", "ServiceRequestError"])
def test_list_subscriptions_raises_on_http_and_transport_failure(monkeypatch, error_name):
    exceptions = pytest.importorskip("azure.core.exceptions")
    error = getattr(exceptions, error_name)(message="AuthorizationFailed: no access")
    monkeypatch.setattr(utils, "get_azure_client", lambda service: _subscription_client_raising(error))
    monkeypatch.setattr(utils, "detect_environment", lambda: "government")

    with pytest.raises(utils.AzureAccessError) as excinfo:
        utils.list_subscriptions()

    assert excinfo.value.__cause__ is error
    assert "AzureUSGovernment" in excinfo.value.hint


def test_list_subscriptions_returns_empty_list_only_when_azure_returns_none(monkeypatch):
    pytest.importorskip("azure.core.exceptions")

    class _Subscriptions:
        def list(self):
            return iter(())

    class _Client:
        subscriptions = _Subscriptions()

    monkeypatch.setattr(utils, "get_azure_client", lambda service: _Client())
    assert utils.list_subscriptions() == []


def test_prompt_menu_treats_eof_as_exit(monkeypatch, capsys):
    monkeypatch.delenv("STRATUSSCAN_AUTO_RUN", raising=False)

    def closed_stdin(prompt=""):
        raise EOFError

    monkeypatch.setattr("builtins.input", closed_stdin)
    assert utils.prompt_menu("T", ["a"], allow_back=False, allow_exit=True) == "exit"
    assert utils.prompt_menu("T", ["a"], allow_back=True, allow_exit=False) == "back"


def test_is_service_available_in_environment_contract():
    assert utils.is_service_available_in_environment("compute", "public") is True
    assert utils.is_service_available_in_environment("network", "government") is True


def test_get_azure_client_rejects_unknown_service():
    try:
        utils.get_azure_client("not-a-real-service", "sub")
    except ValueError as exc:
        assert "not-a-real-service" in str(exc)
    else:
        raise AssertionError("expected ValueError for unknown service key")


def test_client_map_entries_are_well_formed():
    for key, entry in utils._CLIENT_MAP.items():
        assert key == key.lower(), f"{key!r} must be lowercase"
        assert isinstance(entry, tuple) and len(entry) == 3, f"{key!r} must be a 3-tuple"
        module_path, class_name, needs_sub = entry
        assert module_path.startswith("azure."), f"{key!r} module path looks wrong: {module_path}"
        assert class_name and class_name[0].isupper(), f"{key!r} class name looks wrong: {class_name}"
        assert isinstance(needs_sub, bool), f"{key!r} needs_sub must be bool"


class _StrEnum(str, enum.Enum):
    INBOUND = "Inbound"


class _PlainEnum(enum.Enum):
    STANDARD = "Standard"
    COUNT = 3


def _read_rows(path, sheet):
    from openpyxl import load_workbook

    return [[c.value for c in row] for row in load_workbook(path)[sheet].iter_rows(min_row=2)]


def test_s_normalizes_none_enum_and_scalars():
    assert utils.s(None) == ""
    assert utils.s(_StrEnum.INBOUND) == "Inbound"
    assert utils.s(_PlainEnum.STANDARD) == "Standard"
    assert utils.s(_PlainEnum.COUNT) == "3"
    assert utils.s("plain") == "plain"
    assert utils.s(0) == "0"
    assert utils.s(False) == "False"


def test_s_never_renders_the_enum_class_name():
    assert "INBOUND" not in utils.s(_StrEnum.INBOUND)
    assert "_StrEnum" not in utils.s(_StrEnum.INBOUND)


def test_save_dataframe_to_excel_writes_enum_values_and_blank_none(tmp_path):
    import pandas as pd

    df = pd.DataFrame(
        [
            {"Direction": _StrEnum.INBOUND, "Tier": _PlainEnum.STANDARD, "Note": None, "Count": 2},
            {"Direction": "Outbound", "Tier": _PlainEnum.STANDARD, "Note": "x", "Count": 0},
        ]
    )
    path = str(tmp_path / "single.xlsx")

    utils.save_dataframe_to_excel(df, path)

    assert _read_rows(path, "Export") == [
        ["Inbound", "Standard", None, 2],
        ["Outbound", "Standard", "x", 0],
    ]


def test_save_dataframe_to_excel_does_not_mutate_the_callers_dataframe(tmp_path):
    import pandas as pd

    df = pd.DataFrame([{"Direction": _StrEnum.INBOUND}])

    utils.save_dataframe_to_excel(df, str(tmp_path / "keep.xlsx"))

    assert df["Direction"].iloc[0] is _StrEnum.INBOUND


def test_save_multiple_dataframes_to_excel_normalizes_every_sheet(tmp_path):
    import pandas as pd

    path = str(tmp_path / "multi.xlsx")

    utils.save_multiple_dataframes_to_excel(
        {
            "NSGs": pd.DataFrame([{"Direction": _StrEnum.INBOUND, "Rules": 4}]),
            "Rules": pd.DataFrame([{"Access": _PlainEnum.STANDARD, "Note": None}]),
        },
        path,
    )

    assert _read_rows(path, "NSGs") == [["Inbound", 4]]
    assert _read_rows(path, "Rules") == [["Standard", None]]


@pytest.mark.parametrize(
    "value, expected",
    [
        ("=1+1", "'=1+1"),
        ("+SUM(A1)", "+SUM(A1)"),
        ("-2+3", "-2+3"),
        ("@SUM(A1)", "@SUM(A1)"),
        ("==", "'=="),
        ("plain", "plain"),
        ("a=b", "a=b"),
        ("", ""),
        (-1, -1),
        (-1.5, -1.5),
        (0, 0),
        (None, None),
        (False, False),
    ],
)
def test_guard_formula_prefixes_only_dangerous_strings(value, expected):
    assert utils._guard_formula(value) == expected


def test_save_dataframe_to_excel_neutralizes_formula_cells(tmp_path):
    import pandas as pd
    from openpyxl import load_workbook

    df = pd.DataFrame(
        [
            {"Tag": "=HYPERLINK(\"http://evil\",\"x\")", "Count": -1, "Ratio": -1.5, "Note": "-lead"},
            {"Tag": "safe", "Count": 2, "Ratio": 0.5, "Note": "@x"},
        ]
    )
    path = str(tmp_path / "guard.xlsx")

    utils.save_dataframe_to_excel(df, path)

    ws = load_workbook(path)["Export"]
    assert ws["A2"].data_type == "s"
    assert ws["A2"].value == "'=HYPERLINK(\"http://evil\",\"x\")"
    assert ws["A3"].value == "safe"
    assert (ws["B2"].value, ws["B2"].data_type) == (-1, "n")
    assert (ws["C2"].value, ws["C2"].data_type) == (-1.5, "n")
    assert (ws["D2"].value, ws["D2"].data_type) == ("-lead", "s")
    assert (ws["D3"].value, ws["D3"].data_type) == ("@x", "s")
    assert df["Tag"].iloc[0].startswith("=")


def test_save_multiple_dataframes_to_excel_neutralizes_every_sheet(tmp_path):
    import pandas as pd
    from openpyxl import load_workbook

    path = str(tmp_path / "guard-multi.xlsx")
    utils.save_multiple_dataframes_to_excel(
        {"A": pd.DataFrame([{"Name": "=1"}]), "B": pd.DataFrame([{"Name": "+1", "Direction": _StrEnum.INBOUND}])},
        path,
    )
    wb = load_workbook(path)
    assert wb["A"]["A2"].value == "'=1"
    assert (wb["B"]["A2"].value, wb["B"]["A2"].data_type) == ("+1", "s")
    assert wb["B"]["B2"].value == "Inbound"


def test_get_azure_client_reports_renamed_class_as_import_error(monkeypatch):
    monkeypatch.setitem(utils._CLIENT_MAP, "subscription", ("json", "NoSuchClient", False))

    with pytest.raises(ImportError, match="NoSuchClient.*pyproject.toml"):
        utils.get_azure_client("subscription")


def test_get_azure_client_pip_hint_uses_distribution_name_for_nested_modules(monkeypatch):
    monkeypatch.setitem(
        utils._CLIENT_MAP, "resource", ("azure.mgmt.doesnotexist.resources", "Client", True)
    )

    with pytest.raises(ImportError, match=r"^Missing package: pip install azure-mgmt-doesnotexist$"):
        utils.get_azure_client("resource", "sub-id")


# --- SSAZR-114: failure ≠ empty ---------------------------------------------------


def test_no_resources_found_message_and_noun():
    exc = utils.NoResourcesFound("storage accounts")
    assert str(exc) == "No storage accounts found."
    assert exc.noun == "storage accounts"


def test_export_result_defaults_errors_to_a_fresh_list():
    first = utils.ExportResult(rows=1, filename="a.xlsx")
    second = utils.ExportResult(rows=2, filename=None)
    first.errors.append({"Scope": "x"})
    assert second.errors == []


def test_error_record_prefers_service_code_then_status_then_type():
    from types import SimpleNamespace

    exceptions = pytest.importorskip("azure.core.exceptions")
    coded = exceptions.HttpResponseError(message="first line\nsecond line")
    coded.error = SimpleNamespace(code="AuthorizationFailed")
    assert utils.error_record("rg1", "list", coded) == {
        "Scope": "rg1", "Operation": "list", "Error Code": "AuthorizationFailed", "Message": "first line",
    }

    status_only = exceptions.HttpResponseError(message="nope")
    status_only.status_code = 429
    assert utils.error_code(status_only) == "HTTP 429"

    assert utils.error_code(RuntimeError("x")) == "RuntimeError"
    assert utils.error_message(RuntimeError("")) == "RuntimeError"


def test_save_dataframe_to_excel_appends_errors_sheet_only_when_errors(tmp_path):
    pd = pytest.importorskip("pandas")
    load_workbook = pytest.importorskip("openpyxl").load_workbook
    df = pd.DataFrame([{"Name": "a"}])

    clean = tmp_path / "clean.xlsx"
    utils.save_dataframe_to_excel(df, str(clean), sheet_name="Data", errors=[])
    assert load_workbook(clean).sheetnames == ["Data"]

    partial = tmp_path / "partial.xlsx"
    errors = [utils.error_record("acct1", "blob_containers.list", RuntimeError("=boom"))]
    utils.save_dataframe_to_excel(df, str(partial), sheet_name="Data", errors=errors)
    wb = load_workbook(partial)
    assert wb.sheetnames == ["Data", "Errors"]
    ws = wb["Errors"]
    assert [c.value for c in ws[1]] == list(utils.ERROR_COLUMNS)
    assert [c.value for c in ws[2]] == ["acct1", "blob_containers.list", "RuntimeError", "'=boom"]


def test_save_multiple_dataframes_to_excel_appends_errors_sheet_last(tmp_path):
    pd = pytest.importorskip("pandas")
    load_workbook = pytest.importorskip("openpyxl").load_workbook
    sheets = {"Servers": pd.DataFrame([{"S": 1}]), "Databases": pd.DataFrame([{"D": 2}])}
    path = tmp_path / "multi.xlsx"

    utils.save_multiple_dataframes_to_excel(sheets, str(path), errors=[{"Scope": "s1", "Operation": "op", "Error Code": "HTTP 403", "Message": "m"}])

    assert load_workbook(path).sheetnames == ["Servers", "Databases", "Errors"]
    assert list(sheets) == ["Servers", "Databases"]


def test_run_manifest_round_trip_filters_by_run_id(monkeypatch, tmp_path):
    manifest = tmp_path / ".run-manifest.jsonl"
    monkeypatch.setattr(utils, "manifest_path", lambda: manifest)
    monkeypatch.setenv("STRATUSSCAN_RUN_ID", "run-1")
    utils.record_run_result(script="a", status="OK", rows=1)
    monkeypatch.delenv("STRATUSSCAN_RUN_ID")
    utils.record_run_result(script="b", status="EMPTY", rows=0)
    manifest.write_text(manifest.read_text(encoding="utf-8") + "not json\n", encoding="utf-8")
    utils.record_run_result(run_id="run-1", script="c", status="FAILED")

    run_1 = utils.read_run_results("run-1")
    assert [(r["script"], r["status"]) for r in run_1] == [("a", "OK"), ("c", "FAILED")]
    assert all(r["timestamp"] for r in run_1)
    assert [r["script"] for r in utils.read_run_results("standalone")] == ["b"]
    assert utils.read_run_results("nope") == []


def test_read_run_results_without_manifest_is_empty(monkeypatch, tmp_path):
    monkeypatch.setattr(utils, "manifest_path", lambda: tmp_path / "missing.jsonl")
    assert utils.read_run_results("x") == []


def test_archive_outputs_with_run_id_excludes_stale_workbooks(monkeypatch, tmp_path):
    import zipfile

    output_dir = tmp_path / "output"
    output_dir.mkdir()
    monkeypatch.setattr(utils, "__file__", str(tmp_path / "utils.py"))
    this_run = output_dir / "SUB-storage-accounts-all-export-09.18.2026.xlsx"
    stale = output_dir / "OTHER-vms-all-export-01.01.2026.xlsx"
    report = output_dir / "run-report-all-09.18.2026-101010.xlsx"
    gone = output_dir / "SUB-deleted-all-export-09.18.2026.xlsx"
    for f in (this_run, stale, report):
        f.write_text("xlsx")
    monkeypatch.setenv("STRATUSSCAN_RUN_ID", "09.18.2026-101010")
    utils.record_run_result(script="storage-accounts", status="OK", file=str(this_run))
    utils.record_run_result(script="vms", status="EMPTY", file=None)
    utils.record_run_result(script="deleted", status="OK", file=str(gone))

    zip_path = utils.archive_outputs("all", run_id="09.18.2026-101010")

    assert zip_path == str(output_dir / "exports-all-09.18.2026-101010.zip")
    with zipfile.ZipFile(zip_path) as zf:
        assert sorted(zf.namelist()) == sorted([this_run.name, report.name])
    assert utils.archive_outputs("all", run_id="no-such-run") is None


def test_setup_logging_filename_carries_script_sub_and_seconds(monkeypatch, tmp_path):
    import logging
    import re

    monkeypatch.setattr(utils, "__file__", str(tmp_path / "utils.py"))
    monkeypatch.delenv("STRATUSSCAN_SUBSCRIPTION_ID", raising=False)
    try:
        utils.setup_logging("storage-accounts-export", subscription_id="12345678-aaaa-bbbb-cccc-dddddddddddd")
        (explicit,) = (tmp_path / "logs").glob("*.log")
        assert re.fullmatch(
            r"logs-storage-accounts-export-12345678-\d{2}\.\d{2}\.\d{4}-\d{6}\.log", explicit.name
        )
        explicit.unlink()

        monkeypatch.setenv("STRATUSSCAN_SUBSCRIPTION_ID", "abcdef01-0000-0000-0000-000000000000")
        utils.setup_logging("vms-export")
        (from_env,) = (tmp_path / "logs").glob("*.log")
        assert from_env.name.startswith("logs-vms-export-abcdef01-")
        from_env.unlink()

        monkeypatch.delenv("STRATUSSCAN_SUBSCRIPTION_ID")
        utils.setup_logging("vms-export")
        (nosub,) = (tmp_path / "logs").glob("*.log")
        assert nosub.name.startswith("logs-vms-export-nosub-")
    finally:
        for handler in list(logging.getLogger("stratusscan").handlers):
            handler.close()
        logging.getLogger("stratusscan").handlers = []
