#!/usr/bin/env python3
"""
StratusScanCLI-Azure — Shared Utilities Module
Version: v0.1.0

Shared utility functions for all StratusScan exporter scripts.
Handles credential management, client factory, environment detection,
logging, Excel output, and config I/O.

Design constraints:
- No print() calls in this module — return structured results only
- setup_logging() is called from main() (entry points) or runner.run_exporter() (exporters); never on import
- CloudShell-first: all paths must work in a fresh Azure Cloud Shell session
"""

import datetime
import enum
import json
import logging
import os
import platform
import re
import sys
import threading
import warnings
from dataclasses import dataclass, field
from importlib.metadata import PackageNotFoundError
from importlib.metadata import version as _pkg_version
from pathlib import Path
from typing import Any

# ---------------------------------------------------------------------------
# Version
# ---------------------------------------------------------------------------

def get_version() -> str:
    try:
        return _pkg_version("stratusscan-cli-azure")
    except PackageNotFoundError:
        return "dev"


# ---------------------------------------------------------------------------
# Logging
# ---------------------------------------------------------------------------

logger: logging.Logger | None = None
_logging_configured = False


def _cleanup_old_logs(logs_dir: Path, retention_days: int = 14) -> None:
    try:
        cutoff = (datetime.datetime.now() - datetime.timedelta(days=retention_days)).timestamp()
        for f in logs_dir.glob("*.log"):
            try:
                if f.stat().st_mtime < cutoff:
                    f.unlink()
            except Exception:
                pass
    except Exception:
        pass


def _log_subscription_tag(subscription_id: str | None) -> str:
    sub_id = (subscription_id or os.environ.get("STRATUSSCAN_SUBSCRIPTION_ID", "")).strip()
    return sub_id[:8] if sub_id else "nosub"


def setup_logging(
    script_name: str = "stratusscan",
    log_to_file: bool = True,
    subscription_id: str | None = None,
) -> logging.Logger:
    """
    Configure the shared logger. runner.run_exporter() calls this with the resolved
    subscription; without one the file name falls back to the STRATUSSCAN_SUBSCRIPTION_ID
    the orchestrator injects — that is what keeps a multi-subscription Run All from
    overwriting one subscription's log with another's.
    """
    global logger, _logging_configured

    logger = logging.getLogger("stratusscan")
    logger.setLevel(logging.DEBUG)
    logger.handlers = []

    console_fmt = logging.Formatter("%(asctime)s - %(levelname)s - %(message)s", datefmt="%Y-%m-%d %H:%M:%S")
    file_fmt = logging.Formatter(
        "%(asctime)s - %(name)s - %(levelname)s - %(funcName)s:%(lineno)d - %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S",
    )

    console_handler = logging.StreamHandler(sys.stdout)
    console_handler.setLevel(logging.WARNING)
    console_handler.setFormatter(console_fmt)
    logger.addHandler(console_handler)

    if log_to_file:
        try:
            logs_dir = Path(__file__).parent / "logs"
            logs_dir.mkdir(exist_ok=True)
            _cleanup_old_logs(logs_dir)
            timestamp = datetime.datetime.now().strftime("%m.%d.%Y-%H%M%S")
            sub_tag = _log_subscription_tag(subscription_id)
            log_path = logs_dir / f"logs-{script_name}-{sub_tag}-{timestamp}.log"
            fh = logging.FileHandler(log_path, mode="w", encoding="utf-8")
            fh.setLevel(logging.DEBUG)
            fh.setFormatter(file_fmt)
            logger.addHandler(fh)
            logger.info("StratusScan logging initialized — %s", log_path)
        except Exception as exc:
            logger.warning("File logging unavailable: %s", exc)

    for noisy_logger in ("azure", "msrest", "urllib3"):
        logging.getLogger(noisy_logger).setLevel(logging.ERROR)

    _logging_configured = True
    return logger


def set_console_level(level: int) -> None:
    """Raise or lower the console handler only; the file handler stays at DEBUG."""
    for handler in get_logger().handlers:
        if isinstance(handler, logging.StreamHandler) and not isinstance(handler, logging.FileHandler):
            handler.setLevel(level)


def get_logger() -> logging.Logger:
    global logger, _logging_configured
    if logger is None:
        nl = logging.getLogger("stratusscan")
        if not nl.handlers:
            nl.addHandler(logging.NullHandler())
        return nl
    return logger


def log_script_start(script_name: str, description: str = "") -> None:
    log = get_logger()
    log.info("=" * 80)
    log.info("Script: %s  |  %s", script_name, description)
    log.info("Version: %s", get_version())
    log.info("Started: %s", datetime.datetime.now().isoformat())
    log.info("=" * 80)


def log_system_info() -> None:
    log = get_logger()
    log.info("Platform: %s %s", platform.system(), platform.release())
    log.info("Python: %s", sys.version.split()[0])


# ---------------------------------------------------------------------------
# CI / automation mode
# ---------------------------------------------------------------------------

def is_auto_run() -> bool:
    return os.environ.get("STRATUSSCAN_AUTO_RUN", "").strip() == "1"


def get_auto_subscriptions() -> list[str]:
    raw = os.environ.get("STRATUSSCAN_SUBSCRIPTIONS", "").strip()
    if not raw:
        return []
    return [s.strip() for s in raw.split(",") if s.strip()]


# ---------------------------------------------------------------------------
# Timestamp & filename helpers
# ---------------------------------------------------------------------------

def get_current_timestamp() -> str:
    """Return today's date as MM.DD.YYYY."""
    return datetime.datetime.now().strftime("%m.%d.%Y")


def _sanitize_name(name: str) -> str:
    """Replace characters that are unsafe in filenames with hyphens."""
    return re.sub(r"[^\w\-]", "-", name).strip("-")


OUTPUT_DIR_ENV = "STRATUSSCAN_OUTPUT_DIR"


def output_dir() -> Path:
    """
    Return the directory every artifact of a run lands in, creating it if needed.

    Precedence: STRATUSSCAN_OUTPUT_DIR (what runner.py's --output-dir sets), then
    config.json's output_dir, then output/ beside this module. A relative path is
    resolved against the project root so a Cloud Shell `cd` cannot scatter exports.
    Workbooks, the run manifest, the run report and the zip all route through here.
    """
    raw = os.environ.get(OUTPUT_DIR_ENV, "").strip()
    if not raw:
        configured = get_config().get("output_dir")
        raw = configured.strip() if isinstance(configured, str) and configured.strip() else "output"
    path = Path(raw).expanduser()
    if not path.is_absolute():
        path = Path(__file__).parent / path
    path.mkdir(parents=True, exist_ok=True)
    return path


def create_export_filename(subscription_name: str, resource_type: str, suffix: str) -> str:
    """
    Return a full path to the output file.

    Format: {output_dir}/{SUBSCRIPTION-NAME}-{resource-type}-{suffix}-export-{MM.DD.YYYY}.xlsx
    """
    out_dir = output_dir()
    safe_sub = _sanitize_name(subscription_name).upper()
    date_str = get_current_timestamp()
    filename = f"{safe_sub}-{resource_type}-{suffix}-export-{date_str}.xlsx"
    return str(out_dir / filename)


def extract_resource_group(resource_id: str | None) -> str:
    """
    Return the resource group name from an Azure resource ID, or "" if absent.

    Azure APIs are inconsistent about the resourceGroups segment casing, so
    callers should use this instead of a literal split on "/resourceGroups/".
    """
    if not resource_id:
        return ""
    parts = resource_id.split("/")
    for i, part in enumerate(parts):
        if part.lower() == "resourcegroups" and i + 1 < len(parts):
            return parts[i + 1]
    return ""


def list_subscription_wide(operations: Any, *method_names: str) -> Any:
    """
    Return the lazy paged result of the first subscription-wide list method the
    operation group exposes; the caller materializes it with list().

    The azure-mgmt-* packages disagree on what the subscription-wide listing is
    called (list / list_by_subscription / list_all) and rename it across majors.
    Raises AttributeError naming every method tried so a renamed method fails
    loudly instead of degrading into a per-resource-group scan.
    """
    if not method_names:
        raise ValueError("list_subscription_wide needs at least one method name")
    for name in method_names:
        method = getattr(operations, name, None)
        if callable(method):
            return method()
    raise AttributeError(
        f"{type(operations).__name__} has none of the subscription-wide list methods "
        f"tried: {', '.join(method_names)}"
    )


def _run_files(out_dir: Path, run_id: str) -> list[Path]:
    """Workbooks the manifest attributes to run_id, plus that run's report, that still exist."""
    files: dict[str, Path] = {}
    for record in read_run_results(run_id):
        path = record.get("file")
        if path and Path(path).exists():
            files[str(Path(path).resolve())] = Path(path)
    for report in out_dir.glob(f"run-report-*-{run_id}.xlsx"):
        files[str(report.resolve())] = report
    return sorted(files.values(), key=lambda p: p.name)


def archive_outputs(label: str | None = None, run_id: str | None = None) -> str | None:
    """
    Bundle exports from output/ into a single zip.

    With run_id, only the workbooks the run manifest attributes to that run (plus
    the run report) are included — output/ persists across Cloud Shell sessions,
    so a blanket glob would commingle stale runs, other subscriptions and other
    engagements into an evidence package. Without run_id, every .xlsx is bundled.

    Returns the zip path, or None if there are no exports to archive.
    """
    import zipfile

    out_dir = output_dir()
    if run_id:
        exports = _run_files(out_dir, run_id)
        stamp = re.sub(r"[^\w.\-]", "-", run_id).strip("-")
    else:
        exports = sorted(out_dir.glob("*.xlsx"))
        stamp = get_current_timestamp()
    if not exports:
        return None

    label_part = f"-{_sanitize_name(label).lower()}" if label else ""
    zip_path = out_dir / f"exports{label_part}-{stamp}.zip"
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for export in exports:
            zf.write(export, arcname=export.name)
    return str(zip_path)


# ---------------------------------------------------------------------------
# Export outcome contract
# ---------------------------------------------------------------------------

class NoResourcesFound(Exception):  # noqa: N818 — a control-flow signal (exit 3), not an error
    """
    Raised by an exporter's main() when the primary listing returned nothing.

    The argument is the plural noun ("storage accounts"); the runner prints the
    message and exits 3 so an empty inventory is never confused with a failure
    or with a run that never happened.
    """

    def __init__(self, noun: str) -> None:
        super().__init__(f"No {noun} found.")
        self.noun = noun


ERROR_COLUMNS = ("Scope", "Operation", "Error Code", "Message")


@dataclass
class ExportResult:
    """What an exporter's main() hands back to the runner."""

    rows: int
    filename: str | None
    errors: list[dict[str, str]] = field(default_factory=list)


def error_code(exc: BaseException) -> str:
    """Service error code (e.g. AuthorizationFailed), else HTTP status, else the exception type."""
    code = getattr(getattr(exc, "error", None), "code", None)
    if code:
        return s(code)
    status = getattr(exc, "status_code", None)
    if status:
        return f"HTTP {status}"
    return type(exc).__name__


def error_message(exc: BaseException) -> str:
    """First non-empty line of the error text — service messages run to many lines."""
    text = getattr(exc, "message", None) or str(exc)
    for line in str(text).splitlines():
        if line.strip():
            return line.strip()
    return type(exc).__name__


def error_record(scope: str, operation: str, exc: BaseException) -> dict[str, str]:
    """One row for the Errors sheet: which parent failed, on which call, and why."""
    return {
        "Scope": scope,
        "Operation": operation,
        "Error Code": error_code(exc),
        "Message": error_message(exc),
    }


# ---------------------------------------------------------------------------
# Run manifest (one JSON line per exporter outcome)
# ---------------------------------------------------------------------------

_MANIFEST_NAME = ".run-manifest.jsonl"


def get_run_id() -> str:
    return os.environ.get("STRATUSSCAN_RUN_ID", "").strip() or "standalone"


def manifest_path() -> Path:
    return output_dir() / _MANIFEST_NAME


def record_run_result(**fields: Any) -> str | None:
    """
    Append one outcome line to output/.run-manifest.jsonl and return its path.

    Recording must never break an export, so an unwritable manifest is logged
    and swallowed here — the exit code still carries the outcome.
    """
    record: dict[str, Any] = {
        "run_id": get_run_id(),
        "timestamp": datetime.datetime.now().isoformat(timespec="seconds"),
    }
    record.update(fields)
    try:
        path = manifest_path()
        with open(path, "a", encoding="utf-8") as fh:
            fh.write(json.dumps(record, default=str) + "\n")
        return str(path)
    except OSError as exc:
        get_logger().warning("Run manifest not written: %s", exc)
        return None


def read_run_results(run_id: str) -> list[dict[str, Any]]:
    """Return the manifest records for run_id in file order; malformed lines are skipped."""
    path = manifest_path()
    if not path.exists():
        return []
    records: list[dict[str, Any]] = []
    try:
        with open(path, encoding="utf-8") as fh:
            for line in fh:
                line = line.strip()
                if not line:
                    continue
                try:
                    record = json.loads(line)
                except json.JSONDecodeError:
                    continue
                if isinstance(record, dict) and record.get("run_id") == run_id:
                    records.append(record)
    except OSError as exc:
        get_logger().warning("Run manifest unreadable: %s", exc)
    return records


# ---------------------------------------------------------------------------
# Excel output
# ---------------------------------------------------------------------------

def _adjust_column_widths(ws) -> None:
    """Auto-fit column widths in an openpyxl worksheet."""
    try:
        from openpyxl.utils import get_column_letter
        for col_idx, col_cells in enumerate(ws.columns, 1):
            max_len = 0
            for cell in col_cells:
                try:
                    cell_len = len(str(cell.value)) if cell.value is not None else 0
                    max_len = max(max_len, cell_len)
                except Exception:
                    pass
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 60)
    except Exception:
        pass


def s(value: Any) -> str:
    """Return a cell-safe string: None → "", Enum member → its value, else str(value)."""
    if value is None:
        return ""
    if isinstance(value, enum.Enum):
        value = value.value
    return str(value)


def _enum_to_value(value: Any) -> Any:
    return value.value if isinstance(value, enum.Enum) else value


# OWASP "CSV Injection": a cell starting with any of these is evaluated as a
# formula by spreadsheet applications. Tags, names and descriptions come from
# the tenant, so they are attacker-controlled input.
# openpyxl types a str starting with "=" as a formula; every other str stays a
# plain string cell in xlsx, so only "=" needs neutralizing.
_FORMULA_TRIGGERS = ("=",)


def _guard_formula(value: Any) -> Any:
    """Prefix a str cell with ' when it would otherwise be typed as a formula. Non-str values pass through."""
    if isinstance(value, str) and value.startswith(_FORMULA_TRIGGERS):
        return "'" + value
    return value


def _to_cell(value: Any) -> Any:
    return _guard_formula(_enum_to_value(value))


def _normalize_cells(df):
    from pandas.api.types import is_bool_dtype, is_datetime64_any_dtype, is_numeric_dtype

    out = df.copy()
    for col in out.columns:
        dtype = out[col].dtype
        # pandas 3 infers str-mixin Enum columns as StringDtype, so filtering on
        # object dtype would skip exactly the columns that need normalizing.
        if is_numeric_dtype(dtype) or is_bool_dtype(dtype) or is_datetime64_any_dtype(dtype):
            continue
        out[col] = out[col].map(_to_cell)
    return out


def _errors_frame(errors: list[dict[str, str]] | None):
    import pandas as pd

    return pd.DataFrame(list(errors or []), columns=list(ERROR_COLUMNS))


def save_dataframe_to_excel(
    df,
    filename: str,
    sheet_name: str = "Export",
    errors: list[dict[str, str]] | None = None,
) -> str:
    """
    Write a single DataFrame to an Excel workbook.

    When errors is non-empty an "Errors" sheet (Scope, Operation, Error Code,
    Message) is appended so a partial export carries its own gaps.

    Returns the filename on success, raises on failure.
    """
    from openpyxl import load_workbook

    if errors:
        return save_multiple_dataframes_to_excel({sheet_name: df}, filename, errors=errors)

    log = get_logger()
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            _normalize_cells(df).to_excel(
                filename, index=False, sheet_name=sheet_name, engine="openpyxl"
            )
        wb = load_workbook(filename)
        ws = wb[sheet_name]
        _adjust_column_widths(ws)
        wb.save(filename)
        log.info("Saved %d rows → %s", len(df), filename)
        return filename
    except Exception as exc:
        log.error("Failed to save %s: %s", filename, exc)
        raise


def save_multiple_dataframes_to_excel(
    sheets: dict[str, Any],
    filename: str,
    errors: list[dict[str, str]] | None = None,
) -> str:
    """
    Write multiple DataFrames to an Excel workbook, one sheet per key.

    Args:
        sheets: {sheet_name: DataFrame}
        filename: target file path
        errors: per-scope failures; when non-empty an "Errors" sheet is appended last

    Returns the filename on success.
    """
    import pandas as pd
    from openpyxl import load_workbook

    if errors:
        sheets = {**sheets, "Errors": _errors_frame(errors)}

    log = get_logger()
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            with pd.ExcelWriter(filename, engine="openpyxl") as writer:
                for sheet_name, df in sheets.items():
                    safe_name = sheet_name[:31]  # Excel sheet name limit
                    _normalize_cells(df).to_excel(writer, index=False, sheet_name=safe_name)
        wb = load_workbook(filename)
        for ws in wb.worksheets:
            _adjust_column_widths(ws)
        wb.save(filename)
        log.info("Saved %d sheet(s) → %s", len(sheets), filename)
        return filename
    except Exception as exc:
        log.error("Failed to save %s: %s", filename, exc)
        raise


# ---------------------------------------------------------------------------
# Config
# ---------------------------------------------------------------------------

_config_lock = threading.Lock()
_config_cache: dict | None = None

_CONFIG_PATH = Path(__file__).parent / "config.json"

_DEFAULT_CONFIG: dict = {
    "subscriptions": [],
    "default_subscription_id": "",
    "environment": "auto",
    "output_dir": "output",
    "log_retention_days": 14,
}


def get_config() -> dict:
    """Thread-safe singleton loader for config.json."""
    global _config_cache
    with _config_lock:
        if _config_cache is not None:
            return _config_cache
        if _CONFIG_PATH.exists():
            try:
                with open(_CONFIG_PATH, encoding="utf-8") as fh:
                    loaded = json.load(fh)
                _config_cache = {**_DEFAULT_CONFIG, **loaded}
            except Exception as exc:
                get_logger().warning("Could not load config.json: %s — using defaults", exc)
                _config_cache = dict(_DEFAULT_CONFIG)
        else:
            _config_cache = dict(_DEFAULT_CONFIG)
        return _config_cache


def save_config(data: dict) -> None:
    global _config_cache
    with _config_lock:
        with open(_CONFIG_PATH, "w", encoding="utf-8") as fh:
            json.dump(data, fh, indent=2)
        _config_cache = data


def config_path() -> Path:
    return _CONFIG_PATH


def reload_config() -> dict:
    """
    Drop the cached config so the next read picks up a config.json written by
    another process (configure.py runs as a subprocess of stratusscan.py).

    Only the config cache is cleared: reloading the whole module would also reset
    the credential and logger globals, which is what the old importlib.reload did.
    """
    global _config_cache
    with _config_lock:
        _config_cache = None
    return get_config()


# ---------------------------------------------------------------------------
# Azure environment detection
# ---------------------------------------------------------------------------

# Services unavailable in AzureUSGovernment (extend as needed)
_GOV_UNAVAILABLE: list[str] = []


def detect_azure_cloud() -> str | None:
    """
    Best-effort detection of the active Azure cloud from the Azure CLI context.

    Reads the active cloud from the Azure CLI config ([cloud] name = ...), which
    Cloud Shell and `az cloud set` populate. Honors AZURE_CONFIG_DIR. Returns
    'government', 'public', or None if it cannot be determined.
    """
    import configparser

    config_dir = os.environ.get("AZURE_CONFIG_DIR") or str(Path.home() / ".azure")
    try:
        parser = configparser.ConfigParser()
        parser.read(Path(config_dir) / "config")
        name = parser.get("cloud", "name", fallback="").strip().lower()
    except Exception:
        return None
    if name == "azureusgovernment":
        return "government"
    if name == "azurecloud":
        return "public"
    return None


_GOVERNMENT_NAMES = ("azureusgovernment", "government", "usgov")
_PUBLIC_NAMES = ("azurepubliccloud", "public", "azurecloud")


def _explicit_environment(value: Any) -> str:
    """Return 'government' / 'public' when value names a cloud, else ''."""
    if not isinstance(value, str):
        return ""
    name = value.strip().lower()
    if name in _GOVERNMENT_NAMES:
        return "government"
    if name in _PUBLIC_NAMES:
        return "public"
    return ""


def detect_environment() -> str:
    """
    Return 'government' if running in AzureUSGovernment, otherwise 'public'.

    Detection order:
    1. AZURE_ENVIRONMENT env var ('AzureUSGovernment' → 'government')
    2. config.json 'environment' key, when it names a cloud (not 'auto' / unset)
    3. Auto-detected active Azure CLI cloud (Cloud Shell / `az cloud set`)
    4. Default: 'public'

    Raises:
        ValueError: AZURE_ENVIRONMENT is set to a value that names no known cloud.
            Guessing here would send a credential to the wrong cloud's endpoints.
    """
    env_var = os.environ.get("AZURE_ENVIRONMENT", "").strip()
    if env_var:
        explicit = _explicit_environment(env_var)
        if not explicit:
            raise ValueError(
                f"AZURE_ENVIRONMENT='{env_var}' is not a recognized Azure cloud. "
                f"Accepted values (case-insensitive): {', '.join(_GOVERNMENT_NAMES + _PUBLIC_NAMES)}. "
                f"Unset it to auto-detect."
            )
        return explicit

    configured = get_config().get("environment")
    explicit = _explicit_environment(configured)
    if explicit:
        return explicit
    if isinstance(configured, str) and configured.strip().lower() not in ("", "auto"):
        get_logger().warning(
            "config.json environment '%s' is not a recognized cloud — auto-detecting.", configured
        )

    return detect_azure_cloud() or "public"


def is_service_available_in_environment(service: str, environment: str) -> bool:
    """
    Return False if the given service is unavailable in the given environment.

    Args:
        service: lowercase service key (e.g. 'cosmosdb', 'web')
        environment: 'public' or 'government'
    """
    if environment == "government" and service.lower() in _GOV_UNAVAILABLE:
        get_logger().warning("Service '%s' is not available in AzureUSGovernment — skipping.", service)
        return False
    return True


# ---------------------------------------------------------------------------
# Credential & client factory
# ---------------------------------------------------------------------------

_credential_lock = threading.Lock()
_credential_cache: Any | None = None


def _get_credential():
    """Return a cached DefaultAzureCredential, configured for the active environment."""
    global _credential_cache
    with _credential_lock:
        if _credential_cache is not None:
            return _credential_cache
        try:
            from azure.identity import AzureAuthorityHosts, DefaultAzureCredential
        except ImportError as exc:
            raise ImportError("azure-identity is required: pip install azure-identity") from exc

        environment = detect_environment()
        if environment == "government":
            _credential_cache = DefaultAzureCredential(
                authority=AzureAuthorityHosts.AZURE_GOVERNMENT
            )
        else:
            _credential_cache = DefaultAzureCredential()
        return _credential_cache


def get_credential() -> Any:
    """
    Return the shared credential for use with data-plane SDKs (e.g. Key Vault
    keys/secrets/certificates clients) that are not azure-mgmt-* clients and so
    are not covered by get_azure_client().

    The credential is configured for the active environment; Key Vault data-plane
    clients resolve the correct (public or government) token scope via challenge
    authentication, so no per-cloud scope handling is needed here.
    """
    return _get_credential()


# Lazy import map: service_name → (module_path, class_name, needs_subscription_id)
_CLIENT_MAP: dict[str, tuple] = {
    "subscription": ("azure.mgmt.resource.subscriptions", "SubscriptionClient", False),
    "resource": ("azure.mgmt.resource.resources", "ResourceManagementClient", True),
    "compute": ("azure.mgmt.compute", "ComputeManagementClient", True),
    "network": ("azure.mgmt.network", "NetworkManagementClient", True),
    "storage": ("azure.mgmt.storage", "StorageManagementClient", True),
    "keyvault": ("azure.mgmt.keyvault", "KeyVaultManagementClient", True),
    "authorization": ("azure.mgmt.authorization", "AuthorizationManagementClient", True),
    "containerservice": ("azure.mgmt.containerservice", "ContainerServiceClient", True),
    "web": ("azure.mgmt.web", "WebSiteManagementClient", True),
    "sql": ("azure.mgmt.sql", "SqlManagementClient", True),
    "cosmosdb": ("azure.mgmt.cosmosdb", "CosmosDBManagementClient", True),
    "policy": ("azure.mgmt.resource.policy", "PolicyClient", True),
    "locks": ("azure.mgmt.resource.locks", "ManagementLockClient", True),
    "managementgroups": ("azure.mgmt.managementgroups", "ManagementGroupsMgmtClient", False),
    "security": ("azure.mgmt.security", "SecurityCenter", True),
    "advisor": ("azure.mgmt.advisor", "AdvisorManagementClient", True),
    "monitor": ("azure.mgmt.monitor", "MonitorManagementClient", True),
    "loganalytics": ("azure.mgmt.loganalytics", "LogAnalyticsManagementClient", True),
    "costmanagement": ("azure.mgmt.costmanagement", "CostManagementClient", False),
    "postgresql": ("azure.mgmt.postgresqlflexibleservers", "PostgreSQLManagementClient", True),
    "mysql": ("azure.mgmt.mysqlflexibleservers", "MySQLManagementClient", True),
    "redis": ("azure.mgmt.redis", "RedisManagementClient", True),
    "containerregistry": ("azure.mgmt.containerregistry", "ContainerRegistryManagementClient", True),
    "appcontainers": ("azure.mgmt.appcontainers", "ContainerAppsAPIClient", True),
    "containerinstance": ("azure.mgmt.containerinstance", "ContainerInstanceManagementClient", True),
    "dns": ("azure.mgmt.dns", "DnsManagementClient", True),
    "privatedns": ("azure.mgmt.privatedns", "PrivateDnsManagementClient", True),
    "cdn": ("azure.mgmt.cdn", "CdnManagementClient", True),
    "trafficmanager": ("azure.mgmt.trafficmanager", "TrafficManagerManagementClient", True),
    "eventhub": ("azure.mgmt.eventhub", "EventHubManagementClient", True),
    "servicebus": ("azure.mgmt.servicebus", "ServiceBusManagementClient", True),
    "apimanagement": ("azure.mgmt.apimanagement", "ApiManagementClient", True),
    "logic": ("azure.mgmt.logic", "LogicManagementClient", True),
    "msi": ("azure.mgmt.msi", "ManagedServiceIdentityClient", True),
    "policyinsights": ("azure.mgmt.policyinsights", "PolicyInsightsClient", True),
    "automation": ("azure.mgmt.automation", "AutomationClient", True),
    "batch": ("azure.mgmt.batch", "BatchManagementClient", True),
    "recoveryservices": ("azure.mgmt.recoveryservices", "RecoveryServicesClient", True),
    "recoveryservicesbackup": (
        "azure.mgmt.recoveryservicesbackup", "RecoveryServicesBackupClient", True,
    ),
    "cognitiveservices": (
        "azure.mgmt.cognitiveservices", "CognitiveServicesManagementClient", True,
    ),
    "applicationinsights": ("azure.mgmt.applicationinsights", "ApplicationInsightsManagementClient", True),
}

# The Azure Government token audience is not the endpoint host: azure-mgmt-core
# get_arm_endpoints(AZURE_US_GOVERNMENT) and the Azure CLI (activeDirectoryResourceId)
# both pair this endpoint with the management.core audience.
_GOV_BASE_URL = "https://management.usgovcloudapi.net"
_GOV_ARM_SCOPE = "https://management.core.usgovcloudapi.net/.default"

# Azure Government can lag public cloud SDK defaults. Keep overrides targeted
# to services that have been observed failing against the default api-version.
_GOV_API_VERSIONS: dict[str, str] = {
    "storage": "2025-06-01",
    "web": "2025-03-01",
}


def get_azure_client(service_name: str, subscription_id: str | None = None) -> Any:
    """
    Return an instantiated azure-mgmt-* client for the given service.

    Args:
        service_name: key from _CLIENT_MAP (e.g. 'compute', 'network')
        subscription_id: required for all services except 'subscription'

    Raises:
        ValueError: unknown service name or missing subscription_id
        ImportError: required azure-mgmt-* package not installed
    """
    key = service_name.lower()
    if key not in _CLIENT_MAP:
        raise ValueError(
            f"Unknown service '{service_name}'. Known services: {sorted(_CLIENT_MAP)}"
        )

    module_path, class_name, needs_sub = _CLIENT_MAP[key]
    if needs_sub and not subscription_id:
        raise ValueError(f"subscription_id is required for service '{service_name}'")

    pkg = "-".join(module_path.split(".")[:3])
    try:
        import importlib
        mod = importlib.import_module(module_path)
    except ImportError as exc:
        raise ImportError(f"Missing package: pip install {pkg}") from exc
    try:
        cls = getattr(mod, class_name)
    except AttributeError as exc:
        try:
            installed = _pkg_version(pkg)
        except PackageNotFoundError:
            installed = "unknown"
        raise ImportError(
            f"{module_path} has no {class_name}: installed {pkg} {installed} is outside "
            f"the supported range. Install the version range pinned for {pkg} in pyproject.toml."
        ) from exc

    cred = _get_credential()
    environment = detect_environment()

    kwargs: dict[str, Any] = {}
    if environment == "government":
        kwargs["base_url"] = _GOV_BASE_URL
        kwargs["credential_scopes"] = [_GOV_ARM_SCOPE]
        if key in _GOV_API_VERSIONS:
            kwargs["api_version"] = _GOV_API_VERSIONS[key]

    if needs_sub:
        return cls(cred, subscription_id, **kwargs)
    return cls(cred, **kwargs)


# ---------------------------------------------------------------------------
# Subscription helpers
# ---------------------------------------------------------------------------

class AzureAccessError(Exception):
    """
    Azure could not be reached or refused the credential; the SDK error is __cause__.

    Lets the console scripts tell "access failed" apart from "zero subscriptions"
    without importing azure-core themselves.
    """

    def __init__(self, cause: BaseException, environment: str) -> None:
        lines = str(cause).splitlines()
        summary = next((line.strip() for line in lines if line.strip()), "")
        # DefaultAzureCredential's first line is generic; the chain stops at the
        # credential that hard-failed, so the last attempt names the real cause.
        attempts = [line.strip() for line in lines if line.startswith("\t")]
        if attempts:
            summary = f"{summary} Last attempted: {attempts[-1]}"
        super().__init__(f"{type(cause).__name__}: {summary}")
        self.environment = environment

    @property
    def hint(self) -> str:
        label = "AzureUSGovernment" if self.environment == "government" else "AzurePublicCloud"
        return (
            f"Active cloud detected as {label}. If that is wrong, set AZURE_ENVIRONMENT "
            f"or run configure.py; otherwise check your sign-in (`az login`)."
        )


def list_subscriptions() -> list[dict[str, str]]:
    """
    Return all subscriptions accessible to the current credential.

    Returns:
        List of dicts with keys: id, name, state, tenant_id

    Raises:
        AzureAccessError: authentication, authorization or transport failure. An
            empty list therefore always means the credential really sees nothing.
    """
    client = get_azure_client("subscription")
    from azure.core.exceptions import AzureError

    subs = []
    try:
        for sub in client.subscriptions.list():
            subs.append({
                "id": sub.subscription_id,
                "name": sub.display_name or sub.subscription_id,
                "state": str(sub.state) if sub.state else "",
                "tenant_id": sub.tenant_id or "",
            })
    except AzureError as exc:
        # INFO keeps the full SDK message in the log file only; the caller prints the summary.
        get_logger().info("Failed to list subscriptions: %s", exc)
        raise AzureAccessError(exc, detect_environment()) from exc
    return subs


_SUBSCRIPTION_ID_RE = re.compile(
    r"^[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}$", re.IGNORECASE
)


def is_subscription_id(value: str) -> bool:
    """True when value is a GUID — the only shape an Azure subscription ID takes."""
    return bool(_SUBSCRIPTION_ID_RE.match(str(value).strip()))


def get_subscription_name(subscription_id: str) -> str:
    """Return the display name for a subscription ID, or the ID itself on failure."""
    cfg = get_config()
    for sub in cfg.get("subscriptions", []):
        if sub.get("id") == subscription_id:
            return sub.get("name", subscription_id)
    return subscription_id


def resolve_target_subscription() -> tuple:
    """
    Return (subscription_id, subscription_name) for an exporter run.

    Prefers the subscription injected by stratusscan.py via the
    STRATUSSCAN_SUBSCRIPTION_ID / _NAME env vars — this is how multi-subscription
    runs target each subscription in turn. Falls back to the configured default
    when an exporter is run directly. Returns ("", "") when nothing is configured.
    """
    sub_id = os.environ.get("STRATUSSCAN_SUBSCRIPTION_ID", "").strip()
    if sub_id:
        sub_name = os.environ.get("STRATUSSCAN_SUBSCRIPTION_NAME", "").strip()
        return sub_id, (sub_name or get_subscription_name(sub_id))

    cfg = get_config()
    sub_id = cfg.get("default_subscription_id", "")
    if not sub_id:
        return "", ""
    return sub_id, get_subscription_name(sub_id)
